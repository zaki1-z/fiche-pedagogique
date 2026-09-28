import io
import json
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from google import genai
from google.genai import types
import pypdf
from pptx import Presentation

# Configuration de la page
st.set_page_config(
    page_title="Portail Pédagogique - Physique-Chimie Collège",
    page_icon="⚗️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Mot de passe d'accès pour l'enseignant
MOT_DE_PASSE_VALIDE = "Prof2026@"

# Récupération sécurisée et invisible de la clé API depuis Streamlit Secrets
api_key = st.secrets.get("GEMINI_API_KEY", "")

# Authentification
if "authentifie" not in st.session_state:
    st.session_state.authentifie = False

def verifier_mdp():
    if st.session_state.get("mdp_input") == MOT_DE_PASSE_VALIDE:
        st.session_state.authentifie = True
    else:
        st.error("Mot de passe incorrect. Veuillez réessayer.")

if not st.session_state.authentifie:
    st.markdown("""
        <div style="text-align: center; margin-top: 50px; margin-bottom: 25px;">
            <h1 style="color: #17365D;">🔬 Portail Pédagogique - Physique-Chimie</h1>
            <p style="color: #555; font-size: 16px;">Générateur automatisé de fiches conformes aux Orientations Pédagogiques Officielles Marocaines</p>
        </div>
    """, unsafe_allow_html=True)
    
    col1, col2, col3 = st.columns([1, 1.2, 1])
    with col2:
        st.markdown("### 🔒 Accès Enseignant")
        st.text_input("Code d'accès enseignant :", type="password", key="mdp_input", on_change=verifier_mdp)
        st.button("Se connecter ➔", type="primary", use_container_width=True, on_click=verifier_mdp)
    st.stop()

# --- EXTRACTION MULTI-FORMATS ---
def extraire_texte(uploaded_file):
    nom = uploaded_file.name.lower()
    texte = ""
    if nom.endswith(".pdf"):
        reader = pypdf.PdfReader(uploaded_file)
        texte = "\n".join([page.extract_text() or "" for page in reader.pages])
    elif nom.endswith(".docx"):
        doc = Document(uploaded_file)
        texte = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    elif nom.endswith(".pptx"):
        prs = Presentation(uploaded_file)
        diapos = []
        for i, slide in enumerate(prs.slides):
            textes_slide = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            textes_slide.append(paragraph.text.strip())
            if textes_slide:
                diapos.append(f"[Diapo {i+1}]\n" + "\n".join(textes_slide))
        texte = "\n\n".join(diapos)
    elif nom.endswith(".txt"):
        texte = uploaded_file.read().decode("utf-8", errors="ignore")
    return texte

# --- OUTILS DE FORMATAGE WORD CONFORMES AUX EXEMPLES ---
def appliquer_arriere_plan(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))

def configurer_marges_cellule(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def ecrire_cellule(cell, text, gras=False, couleur=(0, 0, 0), taille=9, align=WD_ALIGN_PARAGRAPH.LEFT, fond=None):
    cell.text = ""
    lignes = str(text).strip().split("\n")
    for i, ligne in enumerate(lignes):
        ligne_propre = ligne.strip()
        if not ligne_propre:
            continue
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.1
        
        run = p.add_run(ligne_propre)
        run.bold = gras
        run.font.name = "Calibri"
        run.font.size = Pt(taille)
        run.font.color.rgb = RGBColor(*couleur)
        
    configurer_marges_cellule(cell)
    if fond:
        appliquer_arriere_plan(cell, fond)

# --- GÉNÉRATION DU DOCUMENT DOCX STRICTEMENT IDENTIQUE AUX MODÈLES ---
def generer_document_docx_officiel(data):
    doc = Document()
    
    # Marges fines (1.2 cm) pour tenir sur format standard
    for section in doc.sections:
        section.top_margin = Inches(0.47)
        section.bottom_margin = Inches(0.47)
        section.left_margin = Inches(0.47)
        section.right_margin = Inches(0.47)

    # 1. En-tête : Tableau 3 colonnes (Séance/Niveau | Titre Fiche | Professeur/Leçon)
    t_header = doc.add_table(rows=2, cols=3)
    t_header.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_header.style = 'Table Grid'
    
    col_w_header = [Inches(2.2), Inches(3.1), Inches(2.2)]
    
    ecrire_cellule(t_header.rows[0].cells[0], f"Séance : {data.get('seance', '1/2')}\nNiveau : {data.get('niveau', '3ème AC')}", gras=True, fond="F2F2F2")
    ecrire_cellule(t_header.rows[0].cells[1], "Fiche Pédagogique", gras=True, taille=13, align=WD_ALIGN_PARAGRAPH.CENTER, fond="E8EEF5", couleur=(23, 54, 93))
    ecrire_cellule(t_header.rows[0].cells[2], f"Prof : {data.get('enseignant', 'BOUSHIB Nezha')}", gras=True, align=WD_ALIGN_PARAGRAPH.RIGHT, fond="F2F2F2")
    
    # Ligne 2 : Titre de la leçon centré
    cell_lecon = t_header.rows[1].cells[0].merge(t_header.rows[1].cells[1]).merge(t_header.rows[1].cells[2])
    ecrire_cellule(cell_lecon, f"{data.get('titre_lecon', 'Leçon : Physique-Chimie')}", gras=True, taille=11, align=WD_ALIGN_PARAGRAPH.CENTER, fond="D9E1F2", couleur=(23, 54, 93))

    for row in t_header.rows:
        for idx, w in enumerate(col_w_header):
            if idx < len(row.cells):
                row.cells[idx].width = w

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 2. Tableau Cadrage Pédagogique (Objectifs, Question séance, Prérequis, Concepts)
    t_cadre = doc.add_table(rows=2, cols=2)
    t_cadre.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cadre.style = 'Table Grid'
    t_cadre.rows[0].cells[0].width = Inches(4.5)
    t_cadre.rows[0].cells[1].width = Inches(3.0)
    t_cadre.rows[1].cells[0].width = Inches(4.5)
    t_cadre.rows[1].cells[1].width = Inches(3.0)

    # Cellule Objectifs
    txt_obj = "Objectifs (Connaissances et Capacités) :\n" + "\n".join([f"• {o}" for o in data.get("objectifs", [])])
    ecrire_cellule(t_cadre.rows[0].cells[0], txt_obj)
    
    # Cellule Question de la Séance
    txt_question = f"Question de la Séance :\n{data.get('question_seance', '')}"
    ecrire_cellule(t_cadre.rows[0].cells[1], txt_question, gras=True, fond="F9FBFD")

    # Cellule Prérequis
    txt_prerequis = "Prérequis :\n" + "\n".join([f"- {p}" for p in data.get("prerequis", [])])
    ecrire_cellule(t_cadre.rows[1].cells[0], txt_prerequis)

    # Cellule Concepts
    txt_concepts = "Concepts clés :\n" + "\n".join([f"- {c}" for c in data.get("concepts", [])])
    ecrire_cellule(t_cadre.rows[1].cells[1], txt_concepts)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 3. Tableau Principal des Activités (6 colonnes conforme aux documents modèles)
    activites = data.get("activites", [])
    t_act = doc.add_table(rows=1 + len(activites), cols=6)
    t_act.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_act.style = 'Table Grid'

    col_widths = [Inches(1.8), Inches(1.1), Inches(1.6), Inches(1.6), Inches(0.55), Inches(0.85)]
    entetes = [
        "Bilan de chaque activité\n(Résumé / Notions)",
        "Supports",
        "Activité de l'élève",
        "Activité du prof",
        "Durée",
        "Activités Interactives"
    ]
    for idx, h in enumerate(entetes):
        ecrire_cellule(t_act.rows[0].cells[idx], h, gras=True, taille=8.5, fond="1F4E79", couleur=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

    for idx, act in enumerate(activites):
        row = t_act.rows[1 + idx]
        # Colonne 1 : Bilan / Résumé
        ecrire_cellule(row.cells[0], act.get("bilan_contenu", ""), taille=8.5)
        # Colonne 2 : Supports
        ecrire_cellule(row.cells[1], act.get("supports", ""), taille=8.5)
        # Colonne 3 : Activité de l'élève
        ecrire_cellule(row.cells[2], act.get("activite_eleve", ""), taille=8.5)
        # Colonne 4 : Activité du professeur
        ecrire_cellule(row.cells[3], act.get("activite_prof", ""), taille=8.5)
        # Colonne 5 : Durée
        ecrire_cellule(row.cells[4], act.get("duree", "15 min"), taille=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        # Colonne 6 : Type & Question interactive
        desc_inter = f"{act.get('type_etape', '')}\n\n{act.get('questions_interactives', '')}"
        ecrire_cellule(row.cells[5], desc_inter, gras=True, taille=8, align=WD_ALIGN_PARAGRAPH.CENTER, fond="F2F2F2")

    for row in t_act.rows:
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = w

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 4. Tableau d'Évaluation & Remédiation (Pied de page)
    t_eval = doc.add_table(rows=2, cols=3)
    t_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_eval.style = 'Table Grid'
    
    w_eval = [Inches(2.7), Inches(2.7), Inches(2.1)]
    titres_eval = ["Connaissances évaluables :", "Capacités évaluables :", "Obstacles rencontrés & Remédiation :"]
    for idx, t in enumerate(titres_eval):
        ecrire_cellule(t_eval.rows[0].cells[idx], t, gras=True, fond="D9E1F2", couleur=(23, 54, 93))

    ecrire_cellule(t_eval.rows[1].cells[0], data.get("connaissances_evaluables", ""))
    ecrire_cellule(t_eval.rows[1].cells[1], data.get("capacites_evaluables", ""))
    ecrire_cellule(t_eval.rows[1].cells[2], data.get("obstacles_remediation", "Difficulté de représentation abstraite / Remédiation par modélisation ou simulation."))

    for row in t_eval.rows:
        for idx, w in enumerate(w_eval):
            row.cells[idx].width = w

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

# --- INTERFACE ENSEIGNANT CLAIRE ET ÉPURÉE ---
st.markdown("<h2 style='color:#17365D;'>⚗️ Générateur de Fiches Pédagogiques de Physique-Chimie</h2>", unsafe_allow_html=True)
st.caption("Modèle officiel d'inspection (Collèges) — Conforme aux Orientations et Programmes Annuels du Maroc")

with st.sidebar:
    st.header("📋 Paramètres de la séance")
    ens_nom = st.text_input("Professeur :", value="Mme BOUSHIB Nezha")
    niveau_select = st.selectbox(
        "Niveau scolaire :",
        ["1ère AC (1ère Année Collège)", "2ème AC (2ème Année Collège)", "3ème AC (3ème Année Collège)"],
        index=2
    )
    seance_num = st.selectbox("Séance :", ["Séance 1/2", "Séance 2/2", "Séance 1/3", "Séance 2/3", "Séance 3/3", "Séance Unique (1/1)"], index=0)
    titre_manuel = st.text_input("Intitulé / Leçon :", value="Leçon : Atomes et Ions")
    
    st.divider()
    if st.button("🚪 Se déconnecter", use_container_width=True):
        st.session_state.authentifie = False
        st.rerun()

col_u, col_t = st.columns([1.1, 0.9])

with col_u:
    st.markdown("#### 1. Support de cours officiel")
    fichier_cours = st.file_uploader(
        "Déposer le support de cours (PDF, Word, PowerPoint, Texte) :",
        type=["pdf", "docx", "pptx", "txt"]
    )

with col_t:
    st.markdown("#### 2. Ou coller des remarques didactiques / résumé")
    texte_libre = st.text_area("Notes sur le contenu / activités :", height=135, placeholder="Ex: Insister sur la distinction cation/anion et l'activité documentaire sur les étiquettes d'eau minérale...")

contenu_source = ""
if fichier_cours is not None:
    try:
        contenu_source = extraire_texte(fichier_cours)
        st.success(f"✅ Document '{fichier_cours.name}' analysé avec succès !")
    except Exception as e:
        st.error(f"Erreur de lecture du document : {e}")
elif texte_libre.strip():
    contenu_source = texte_libre.strip()

st.divider()

if st.button("🚀 Générer la Fiche Pédagogique Officielle (.DOCX)", type="primary", use_container_width=True):
    if not api_key:
        st.error("Clé API non configurée. Veuillez vérifier les Secrets dans Streamlit.")
    elif not contenu_source.strip():
        st.warning("Veuillez fournir un support de cours (fichier ou texte) avant de lancer la génération.")
    else:
        with st.spinner("Conception de la fiche selon le modèle officiel et les instructions du programme annuel..."):
            try:
                client = genai.Client(api_key=api_key)
                
                prompt = f"""
                Tu es un inspecteur pédagogique de l'enseignement secondaire collégial marocain en Physique-Chimie.
                Tu dois générer une fiche pédagogique Word rigoureuse, exactement identique aux fiches modèles d'inspection du Maroc.

                NIVEAU CHOISI : {niveau_select}
                SÉANCE : {seance_num}
                INTITULÉ DE LA LEÇON : {titre_manuel}

                RÈGLES DIDACTIQUES ET PÉDAGOGIQUES DU PROGRAMME :
                1. Respecte scrupuleusement les Orientations Pédagogiques officielles du Ministère :
                   - 1AC : Matière & Environnement, Électricité (circuit simple, dipôles, lois des nœuds/tensions).
                   - 2AC : Matière & Environnement (air, molécules, atomes, réactions chimiques), Lumière et Optique (propagation, lentilles, dispersion), Électricité (courant alternatif, installation domestique).
                   - 3AC : Matériaux (matière/objets, atomes et ions, réactions avec l'air et les solutions pH), Mécanique (mouvement, repos, vitesse, actions mécaniques, forces, équilibre, poids/masse), Électricité (loi d'Ohm, puissance, énergie).
                2. Structure de la séance obligatoire en 3 étapes :
                   - "Activité Introductive" (10 min) : Rappel des prérequis, question de la séance (situation-problème), formulation des hypothèses par les élèves.
                   - "Activité Constructive" (30 min) : Investigation, activités documentaires ou expérimentales (avec matériel précis de labo de collège), manipulation, déductions et calculs des élèves, institutionnalisation intermédiaire.
                   - "BILAN" (20 min) : Synthèse collective, résumé structuré à copier, exercices d'application immédiate.

                Réponds STRICTEMENT par un objet JSON valide suivant exactement cette structure :
                {{
                  "titre_lecon": "{titre_manuel}",
                  "question_seance": "Formulation claire de la question-problème de départ",
                  "objectifs": [
                    "Connaître...",
                    "Savoir écrire...",
                    "Distinguer entre..."
                  ],
                  "prerequis": [
                    "Prérequis 1",
                    "Prérequis 2",
                    "Prérequis 3"
                  ],
                  "concepts": [
                    "Concept 1",
                    "Concept 2"
                  ],
                  "activites": [
                    {{
                      "type_etape": "Activité Introductive",
                      "duree": "10 min",
                      "bilan_contenu": "• Poser la question de la séance\\n• Comprendre le problème posé\\n• Proposer des hypothèses",
                      "supports": "- Documents du manuel\\n- Exemples du quotidien\\n- Tableau",
                      "activite_eleve": "- Répondre aux questions et vérifier ses prérequis.\\n- Lire et s'approprier la situation.\\n- Formuler des hypothèses.",
                      "activite_prof": "- Poser les questions de réactivation.\\n- Écrire la situation-problème au tableau.\\n- Recueillir et noter les hypothèses des élèves.",
                      "questions_interactives": "Discussion ouverte avec les élèves sur la situation de départ."
                    }},
                    {{
                      "type_etape": "Activité Constructive",
                      "duree": "30 min",
                      "bilan_contenu": "Résumé du contenu notionnel construit :\\nI- Définitions et règles...\\n- Démonstrations ou résultats d'expériences...",
                      "supports": "- Matériel de laboratoire (éprouvettes, multimètre, etc.)\\n- Fiche d'activité documentaire\\n- Tableau",
                      "activite_eleve": "- Réaliser l'expérience ou analyser le document.\\n- Interpréter les résultats et répondre aux consignes.\\n- Déduire la règle ou la loi physique/chimique.",
                      "activite_prof": "- Guider l'investigation sans donner la solution.\\n- Poser les questions de guidage.\\n- Superviser les mesures et manipulations expérimentales.",
                      "questions_interactives": "Questions clés guidant la démarche d'investigation."
                    }},
                    {{
                      "type_etape": "BILAN",
                      "duree": "20 min",
                      "bilan_contenu": "Synthèse et Institutionnalisation :\\n- Retenir l'essentiel du cours.\\n- Exercice d'application résolu.",
                      "supports": "- Tableau\\n- Manuel scolaire / Cahier de cours",
                      "activite_eleve": "- Participer à l'élaboration de la synthèse.\\n- Noter le cours sur le cahier.\\n- Résoudre l'exercice d'évaluation formative.",
                      "activite_prof": "- Structurer la réponse finale à la question de départ.\\n- Dicter/noter le résumé institutionnel.\\n- Proposer l'exercice d'évaluation.",
                      "questions_interactives": "Évaluation formative et bilan des acquis."
                    }}
                  ],
                  "connaissances_evaluables": "Connaissances clés à évaluer lors du contrôle...",
                  "capacites_evaluables": "Capacités méthodologiques et d'analyse évaluables...",
                  "obstacles_remediation": "Obstacle didactique prévisible et remédiation pédagogique proposée."
                }}

                Support de cours à traiter :
                \"\"\"{contenu_source[:15000]}\"\"\"
                """

                reponse = client.models.generate_content(
                    model="gemini-2.5-flash",
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        temperature=0.2
                    )
                )

                fiche_data = json.loads(reponse.text)
                
                # Injection des métadonnées du professeur
                fiche_data["enseignant"] = ens_nom
                fiche_data["niveau"] = niveau_select.split(" ")[0]
                fiche_data["seance"] = seance_num
                if not fiche_data.get("titre_lecon"):
                    fiche_data["titre_lecon"] = titre_manuel

                doc_docx = generer_document_docx_officiel(fiche_data)

                st.success("🎉 Fiche pédagogique générée avec succès selon le modèle officiel !")

                nom_fichier = f"{fiche_data['niveau']}_{fiche_data['seance'].replace('/', '-')}_{fiche_data['titre_lecon'].replace(' ', '_')}.docx"
                
                st.download_button(
                    label="📥 Télécharger la Fiche Officielle Word (.DOCX)",
                    data=doc_docx,
                    file_name=nom_fichier,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary"
                )

                with st.expander("👁️ Prévisualiser les éléments de la fiche générée"):
                    st.write(f"**Question de départ :** {fiche_data.get('question_seance')}")
                    st.write("**Objectifs :**", fiche_data.get("objectifs"))
                    st.write("**Prérequis :**", fiche_data.get("prerequis"))
                    st.write("**Concepts :**", fiche_data.get("concepts"))

            except Exception as e:
                st.error(f"Une erreur est survenue pendant la génération : {e}")
